"""Build the downloadable teaching pack; run with the bundled artifact Python."""
from pathlib import Path
import json, re, html, urllib.request, zipfile
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'public' / 'case-resources'
QA = ROOT / 'tmp' / 'worksheet-build'
OUT.mkdir(parents=True, exist_ok=True)
QA.mkdir(parents=True, exist_ok=True)
manifest = []
def register(case, filename, label, note=''):
    manifest.append(dict(case=case, file=filename, name=label, note=note))
def sql(case, filename, text, label):
    (OUT / filename).write_text('-- Database City | MySQL 8.0+ | UTF-8\n-- ข้อมูลสมมติสำหรับฝึกในเครื่องเท่านั้น\n'+text.strip()+'\n', encoding='utf-8')
    register(case, filename, label)
def setup(case):
    return f'''-- สร้างหรือเริ่มข้อมูลฝึกใหม่เฉพาะ db_city_case{case:02d}
-- รันทั้งไฟล์จะล้างตารางฝึกใน schema นี้ ให้สำรองงานที่ทำแล้วก่อน
CREATE DATABASE IF NOT EXISTS db_city_case{case:02d} CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci;
USE db_city_case{case:02d};
'''

sql(5, 'Case05_DDL_Template.sql', '''-- แม่แบบนี้มีเฉพาะคอมเมนต์ ให้เติมคำสั่งเองก่อนรัน
-- 1. เขียนคำสั่งสร้างฐานข้อมูลชื่อ db_city
-- [พิมพ์โค้ดตรงนี้]

-- 2. เลือกใช้งานฐานข้อมูล db_city
-- [พิมพ์โค้ดตรงนี้]

-- 3. สร้าง tb_users มี id เป็น INT และ username เป็น VARCHAR(50)
-- [พิมพ์โค้ดตรงนี้]

-- 4. ตรวจสอบด้วย SHOW TABLES และ DESCRIBE tb_users
-- [พิมพ์โค้ดตรงนี้]
''', 'แม่แบบ SQL สร้างฐานข้อมูลและตาราง')
sql(6, 'Case06_Script_Template.sql', setup(6)+'''-- แม่แบบตั้งใจเว้นช่อง ___ ให้เติมก่อนรัน ไม่ใช่สคริปต์สำเร็จรูป
-- เติมชนิดข้อมูล คีย์หลัก และคีย์นอกตามเงื่อนไข
-- CustomerID เป็นเลขจำนวนเต็ม, CustomerName ยาวไม่เกิน 100 ตัวอักษร
-- Phone ต้องรักษาเลขศูนย์นำหน้า, OrderDate เก็บวันที่, TotalAmount มีทศนิยม 2 ตำแหน่ง
CREATE TABLE Customers (
  CustomerID ___,
  CustomerName ___ NOT NULL,
  Phone ___,
  PRIMARY KEY (___)
);
CREATE TABLE Orders (
  OrderID ___,
  CustomerID ___ NOT NULL,
  OrderDate ___,
  TotalAmount ___,
  PRIMARY KEY (___),
  FOREIGN KEY (___) REFERENCES ___(___)
);
-- ทดสอบ: เพิ่มลูกค้า 1 คน และใบสั่งซื้อที่อ้างอิงลูกค้านั้น
-- ทดลองเพิ่ม Orders ที่ CustomerID = 999 ซึ่งไม่มีใน Customers
-- บันทึกผลว่าฐานข้อมูลปฏิเสธข้อมูลอ้างอิงผิดอย่างไร
''', 'แม่แบบเติมชนิดข้อมูลและ PK FK')
products = '''CREATE TABLE Products (
 ProductID INT PRIMARY KEY, ProductName VARCHAR(100) NOT NULL,
 Price DECIMAL(10,2) NOT NULL, SalesCount INT NOT NULL DEFAULT 0
);
INSERT INTO Products VALUES
(1,'Notebook',120.00,8),(2,'Mouse',350.00,12),(3,'Keyboard',1000.00,6),
(4,'USB Cable',150.00,0),(5,'Headset',800.00,0);
'''
sql(7, 'Case07_DML_Starter.sql', setup(7)+'DROP TABLE IF EXISTS Products;\n'+products+'SELECT * FROM Products ORDER BY ProductID;', 'ข้อมูลสินค้า 5 รายการสำหรับ INSERT UPDATE DELETE')
sql(8, 'Case08_Master_Script.sql', setup(8)+'''DROP TABLE IF EXISTS OrderItems;
DROP TABLE IF EXISTS Orders;
DROP TABLE IF EXISTS Products;
DROP TABLE IF EXISTS Customers;
CREATE TABLE Customers (CustomerID INT PRIMARY KEY, CustomerName VARCHAR(100) NOT NULL);
CREATE TABLE Products (ProductID INT PRIMARY KEY, ProductName VARCHAR(100) NOT NULL, Price DECIMAL(10,2));
CREATE TABLE Orders (OrderID INT PRIMARY KEY, CustomerID INT NOT NULL, OrderDate DATE NOT NULL);
CREATE TABLE OrderItems (OrderID INT NOT NULL, ProductID INT NOT NULL, Quantity INT NOT NULL,
 UnitPrice DECIMAL(10,2) NOT NULL, PRIMARY KEY(OrderID,ProductID));
-- ยังไม่ประกาศ FK เพื่อให้ผู้เรียนเชื่อมในแบบจำลอง EER ด้วยตนเอง
-- Orders.CustomerID -> Customers.CustomerID
-- OrderItems.OrderID -> Orders.OrderID
-- OrderItems.ProductID -> Products.ProductID
INSERT INTO Customers VALUES (1,'Suda'),(2,'Somchai'),(3,'Mali');
INSERT INTO Products VALUES (1,'Coffee',65),(2,'Tea',50),(3,'Cocoa',60);
INSERT INTO Orders VALUES (101,1,'2026-11-01'),(102,2,'2026-11-02');
INSERT INTO OrderItems VALUES (101,1,2,65),(101,2,1,50),(102,3,1,60);
''', 'ข้อมูล 4 ตารางสำหรับฝึกสร้างความสัมพันธ์ EER')

employees = [(i, ('S' if i%5==0 else 'Employee')+f'{i:02d}', f'Family{i:02d}', ['IT','Sales','HR','Finance'][i%4], 18000+(i%12)*2500) for i in range(1,51)]
def values(rows):
    return ',\n'.join('('+','.join("'"+str(v).replace("'","''")+"'" if isinstance(v,str) else str(v) for v in r)+')' for r in rows)+';\n'
emp_sql='''CREATE TABLE Employees (EmployeeID INT PRIMARY KEY, FirstName VARCHAR(50), LastName VARCHAR(50), Department VARCHAR(30), Salary DECIMAL(10,2));
INSERT INTO Employees VALUES\n'''+values(employees)
customers = [(1,'Suda'),(2,'Somchai'),(3,'Mali'),(4,'Siri'),(5,'Anan'),(6,'Nida')]
cust_sql='CREATE TABLE Customers (CustomerID INT PRIMARY KEY, CustomerName VARCHAR(100));\nINSERT INTO Customers VALUES\n'+values(customers)
prod100='CREATE TABLE Products (ProductID INT PRIMARY KEY, ProductName VARCHAR(100), Price DECIMAL(10,2));\nINSERT INTO Products VALUES\n'+values([(i,f'Product {i:03d}',50+(i%20)*50) for i in range(1,101)])
sql(9,'Case09_Query_Practice.sql',setup(9)+'DROP TABLE IF EXISTS Employees;\nDROP TABLE IF EXISTS Customers;\nDROP TABLE IF EXISTS Products;\n'+emp_sql+cust_sql+prod100,'ข้อมูลพนักงาน 50 คนและสินค้า 100 รายการสำหรับ SELECT WHERE')
sales_data=[(i,['North','Central','South'][i%3],f'2026-12-{i:02d}',1000+i*175) for i in range(1,25)]
sales_sql='CREATE TABLE Sales (SaleID INT PRIMARY KEY, Branch VARCHAR(30), SaleDate DATE, Amount DECIMAL(12,2));\nINSERT INTO Sales VALUES\n'+values(sales_data)
sql(10,'Case10_Aggregate_Data.sql',setup(10)+'DROP TABLE IF EXISTS Sales;\nDROP TABLE IF EXISTS Employees;\n'+sales_sql+emp_sql,'ข้อมูลยอดขายและพนักงานสำหรับ SUM AVG COUNT GROUP BY')

def commerce(case):
    return setup(case)+'''DROP VIEW IF EXISTS vw_sales_summary;
DROP PROCEDURE IF EXISTS GetProductPrice;
DROP TABLE IF EXISTS OrderItems;
DROP TABLE IF EXISTS Orders;
DROP TABLE IF EXISTS ProductDeleteLog;
DROP TABLE IF EXISTS Products;
DROP TABLE IF EXISTS Customers;
DROP TABLE IF EXISTS Employees;
DROP TABLE IF EXISTS Sales;
'''+cust_sql+'''CREATE TABLE Products (ProductID INT PRIMARY KEY, ProductName VARCHAR(100), Price DECIMAL(10,2));
INSERT INTO Products VALUES (1,'Coffee',65),(2,'Tea',50),(3,'Cocoa',60),(4,'Cake',95),(5,'Unused Menu',120);
CREATE TABLE Orders (OrderID INT PRIMARY KEY, CustomerID INT NOT NULL, OrderDate DATE,
 FOREIGN KEY(CustomerID) REFERENCES Customers(CustomerID));
INSERT INTO Orders VALUES (101,1,'2026-12-01'),(102,2,'2026-12-02'),(103,1,'2026-12-03'),(104,3,'2026-12-04');
CREATE TABLE OrderItems (OrderID INT, ProductID INT, Quantity INT NOT NULL, UnitPrice DECIMAL(10,2) NOT NULL,
 PRIMARY KEY(OrderID,ProductID), FOREIGN KEY(OrderID) REFERENCES Orders(OrderID), FOREIGN KEY(ProductID) REFERENCES Products(ProductID));
INSERT INTO OrderItems VALUES (101,1,2,65),(101,2,1,50),(102,3,3,60),(103,4,2,95),(104,1,1,65);
-- Customers 4, 5, 6 ไม่มีคำสั่งซื้อ เพื่อทดสอบ LEFT JOIN
'''+emp_sql+sales_sql
sql(11,'Case11_Join_Scenarios.sql',commerce(11),'ฐานข้อมูลลูกค้าและการสั่งซื้อสำหรับ INNER JOIN LEFT JOIN')
sql(12,'Case12_Subquery_Exercises.sql',commerce(12)+'''-- โครงสร้างและข้อมูลอิง Case11 แต่แยก schema เพื่อไม่แก้ชุดเดิม
INSERT INTO Orders VALUES (105,2,'2026-12-05');
INSERT INTO OrderItems VALUES (105,1,5,65),(105,4,3,95);
-- ฝึกเขียน Subquery หาพนักงานที่ Salary สูงกว่าค่าเฉลี่ยทั้งบริษัท
-- ฝึก IN และ EXISTS หาลูกค้าที่มีรายการสั่งซื้อ
''','ข้อมูลพร้อมโจทย์สำหรับ Subquery และ Nested Queries')
sql(13,'Case13_Performance_Script.sql',commerce(13)+'''CREATE VIEW vw_sales_summary AS
 SELECT c.CustomerID, c.CustomerName, SUM(oi.Quantity * oi.UnitPrice) AS TotalAmount
 FROM Customers c JOIN Orders o ON o.CustomerID=c.CustomerID
 JOIN OrderItems oi ON oi.OrderID=o.OrderID
 GROUP BY c.CustomerID,c.CustomerName;
SELECT * FROM vw_sales_summary;
EXPLAIN SELECT * FROM Customers WHERE CustomerName='Suda';
CREATE INDEX idx_customer_name ON Customers(CustomerName);
EXPLAIN SELECT * FROM Customers WHERE CustomerName='Suda';
-- ชุดข้อมูลเล็กอาจยังใช้ full table scan ได้ เป็นการตัดสินใจของ optimizer
-- เปรียบเทียบแผน ไม่สรุปว่าดัชนีทำให้ทุกคำสั่งเร็วขึ้น
''','ตัวอย่าง CREATE VIEW CREATE INDEX และ EXPLAIN')
sql(14,'Case14_Procedures_Triggers.sql',commerce(14)+'''CREATE TABLE ProductDeleteLog (
 LogID INT AUTO_INCREMENT PRIMARY KEY, ProductID INT, ProductName VARCHAR(100), DeletedAt DATETIME NOT NULL
);
DELIMITER //
CREATE PROCEDURE GetProductPrice(IN p_ProductID INT)
BEGIN
 SELECT ProductID, ProductName, Price FROM Products WHERE ProductID=p_ProductID;
END //
CREATE TRIGGER trg_product_deleted AFTER DELETE ON Products FOR EACH ROW
BEGIN
 INSERT INTO ProductDeleteLog(ProductID,ProductName,DeletedAt)
 VALUES(OLD.ProductID,OLD.ProductName,NOW());
END //
DELIMITER ;
CALL GetProductPrice(1);
-- ProductID 5 ไม่มี OrderItems อ้างอิง จึงใช้ทดสอบลบได้
DELETE FROM Products WHERE ProductID=5;
SELECT * FROM ProductDeleteLog;
''','ตัวอย่าง Procedure รับพารามิเตอร์และ Trigger บันทึก Log')

def set_font(run, size=12, bold=False):
    run.font.name='Tahoma'; run.font.size=Pt(size); run.bold=bold
    props=run._element.get_or_add_rPr(); fonts=props.rFonts
    if fonts is not None:
        fonts.set(qn('w:eastAsia'),'Tahoma'); fonts.set(qn('w:cs'),'Tahoma')
    run.font.color.rgb=RGBColor(0,0,0)
def doc_base(title,case):
    doc=Document(); s=doc.sections[0]; s.page_height=Cm(29.7); s.page_width=Cm(21)
    s.top_margin=s.bottom_margin=Cm(1.7); s.left_margin=s.right_margin=Cm(1.8)
    for name in ['Normal','Title','Heading 1','Heading 2']:
        st=doc.styles[name]; st.font.name='Tahoma'; st.font.color.rgb=RGBColor(0,0,0)
        st.font.size=Pt(12 if name=='Normal' else 17 if name=='Title' else 13)
        st.paragraph_format.space_after=Pt(7); st.paragraph_format.line_spacing=1.25
    p=doc.add_paragraph(style='Title');set_font(p.add_run(title),17,True)
    week=case if case<=8 else case+1
    p=doc.add_paragraph(f'รายวิชา 20901-2001 การจัดการฐานข้อมูลเบื้องต้น   หน่วยที่ {case}   สัปดาห์ที่ {week}')
    for r in p.runs:set_font(r,10)
    doc.add_paragraph('ชื่อสกุล ________________________ รหัสนักเรียน ______________ ชั้น ______')
    footer=s.footer.paragraphs[0];set_font(footer.add_run('DATABASE CITY  |  ชุดใบงาน MySQL  |  '),9)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    return doc
def ptext(doc,text,bold=False):
    p=doc.add_paragraph();set_font(p.add_run(text),12,bold);return p
def answer(doc,prompt,lines=3):
    ptext(doc,prompt,True)
    for _ in range(lines):ptext(doc,'________________________________________________________________________')
def image_box(doc,text='วางภาพหน้าจอและเขียนคำอธิบาย',height=5):
    ptext(doc,text,True)
    t=doc.add_table(rows=1,cols=1);t.autofit=False;t.columns[0].width=Cm(17)
    row=t.rows[0]; row.height=Cm(height)
    cell=t.cell(0,0);cell.text='\nพื้นที่วางภาพหลักฐาน\n'
    tcPr=cell._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
    for edge in ['top','left','bottom','right']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'6');e.set(qn('w:color'),'D9D9D9');b.append(e)
    tcPr.append(b)
def save_doc(case,name,doc,label):
    for tree in [doc.styles.element,doc._element]:
        for border in list(tree.xpath('.//w:pBdr')):
            border.getparent().remove(border)
    for table in doc.tables:
        table.style='Table Grid'
        props=table._tbl.tblPr
        borders=props.find(qn('w:tblBorders'))
        if borders is not None:props.remove(borders)
        borders=OxmlElement('w:tblBorders')
        for edge in ['top','left','bottom','right','insideH','insideV']:
            element=OxmlElement('w:'+edge)
            for key,value in [('val','single'),('sz','6'),('color','D9D9D9')]:element.set(qn('w:'+key),value)
            borders.append(element)
        props.append(borders)
        if len(table.rows)>1:
            for cell in table.rows[0].cells:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'163B56');cell._tc.get_or_add_tcPr().append(shade)
                for para in cell.paragraphs:
                    for run in para.runs:set_font(run,10,True);run.font.color.rgb=RGBColor(255,255,255)
            for row in table.rows[1:]:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        for run in para.runs:set_font(run,10)
    doc.save(OUT/name);register(case,name,label)

d=doc_base('ใบงาน 01 วิเคราะห์ปัญหาระบบไฟล์ข้อมูล',1)
ptext(d,'เปิด Case01_Spreadsheet_Sample.xlsx ข้อมูลทุกชื่อและรายการเป็นข้อมูลสมมติ สำรวจทั้ง 3 ชีต แล้วตอบคำถามต่อไปนี้')
answer(d,'1 ระบุจุดข้อมูลซ้ำและข้อมูลไม่สอดคล้องอย่างน้อย 3 จุด พร้อมบอกชื่อชีตและหมายเลขแถว',2)
image_box(d,'แนบภาพจุดที่พบปัญหา แล้ววงหรือทำเครื่องหมายกำกับ',4)
answer(d,'2 หากปรับราคาสินค้าในชีต Products ด้วยมือ จะเกิดปัญหาใดเมื่อแก้ไขไม่ครบทุกบรรทัด',3)
answer(d,'3 สรุปความแตกต่างของ File System และ Database จำนวน 3 ข้อ',3)
ptext(d,'ตรวจเพิ่มเติม รหัสลูกค้าใน Orders ทุกค่าพบใน Customer หรือไม่ อธิบายหนึ่งตัวอย่าง',True)
save_doc(1,'Worksheet_01.docx',d,'ใบงานวิเคราะห์ปัญหาข้อมูล 3 ข้อพร้อมพื้นที่แนบภาพ')
d=doc_base('ใบงาน 02 ระบุคำศัพท์ฐานข้อมูลเชิงสัมพันธ์',2)
ptext(d,'ศึกษาตารางนักศึกษาต่อไปนี้ แล้วระบุ Entity, Attribute และ Primary Key พร้อมเหตุผล')
t=d.add_table(rows=1,cols=4); t.style='Light Shading Accent 1'
for c,v in zip(t.rows[0].cells,['StudentID','FullName','Classroom','Phone']):c.text=v
for row in [('67001','มาลี ใจดี','1/1','0812340001'),('67002','สมชาย สุขใจ','1/1','0812340002'),('67003','มาลี ใจดี','1/2','0812340003')]:
    for c,v in zip(t.add_row().cells,row):c.text=v
answer(d,'1 Entity name คืออะไร',2);answer(d,'2 ระบุ Attribute ทั้ง 4 ฟิลด์',3)
answer(d,'3 ฟิลด์ใดเหมาะเป็น Primary Key และเพราะเหตุใด ชื่อบุคคลซ้ำกันได้หรือไม่',4)
save_doc(2,'Worksheet_02.docx',d,'ใบงานเติมคำ Entity Attribute Primary Key')
d=doc_base('ใบงาน 04 ตรวจการเชื่อมต่อ MySQL Workbench',4)
ptext(d,'เปิด Local Connection แล้วรัน SELECT VERSION(), CURRENT_USER(); ให้ได้ผลลัพธ์ก่อนบันทึกภาพ')
ptext(d,'Capture หน้าจอ MySQL Workbench ที่แสดงการเชื่อมต่อสำเร็จ พร้อมผลคำสั่ง แล้ววางในกรอบ ห้ามแนบรหัสผ่านจริง',True)
image_box(d,height=10);answer(d,'ระบุ Host และ Port ที่ใช้ พร้อมเวอร์ชัน Server',2)
save_doc(4,'Worksheet_04.docx',d,'ใบงานแนบภาพการเชื่อมต่อ Workbench สำเร็จ')
worksheets={
 7:('จัดการข้อมูลด้วย INSERT UPDATE DELETE','Case07_DML_Starter.sql',[
 'Mission 1 เพิ่มสินค้าใหม่ 2 รายการ ใช้ ProductID 6 และ 7 พร้อมกำหนด Price และ SalesCount',
 'Mission 2 ลดราคา ProductID 3 ลง 50 เปอร์เซ็นต์จากราคาตั้งต้น เขียน SELECT ตรวจสอบก่อนและหลัง',
 'Mission 3 ลบสินค้าที่ SalesCount = 0 โดยตรวจแถวด้วย SELECT ก่อนลบ และระบุจำนวนแถวที่ลบ']),
 9:('สืบค้นด้วย SELECT และ WHERE','Case09_Query_Practice.sql',[
 '1 แสดง FirstName และ LastName ของพนักงานทั้งหมด',
 '2 หาพนักงานที่ Salary มากกว่า 25000',
 "3 หาลูกค้าที่ CustomerName ขึ้นต้นด้วย S ใช้ LIKE 'S%'",
 '4 หาสินค้าที่ Price อยู่ระหว่าง 100 ถึง 500 โดยรวมค่าขอบเขตด้วย BETWEEN',
 '5 แสดงพนักงานแผนก IT เรียง Salary จากมากไปน้อย']),
 10:('คำนวณและจัดกลุ่มข้อมูล','Case10_Aggregate_Data.sql',[
 '1 หาผลรวมยอดขายทั้งหมดจาก Sales.Amount ด้วย SUM',
 '2 หาค่าเฉลี่ยยอดขายต่อรายการจาก Sales.Amount ด้วย AVG',
 "3 นับพนักงานใน Employees ที่ Department = 'IT' ด้วย COUNT",
 '4 แสดงยอดขายรวมแยกตาม Sales.Branch ด้วย GROUP BY']),
 11:('เชื่อมตารางด้วย INNER JOIN และ LEFT JOIN','Case11_Join_Scenarios.sql',[
 "1 ใช้ INNER JOIN เชื่อม Customers, Orders, OrderItems และ Products เพื่อแสดงว่า Suda ซื้อสินค้าอะไร จำนวนเท่าใด ใน OrderID ใด",
 '2 ใช้ Customers LEFT JOIN Orders และเงื่อนไข OrderID IS NULL เพื่อหาลูกค้าที่ยังไม่เคยสั่งซื้อ อธิบายว่าต่างจาก INNER JOIN อย่างไร']),
 12:('ค้นหาข้อมูลด้วย Subquery','Case12_Subquery_Exercises.sql',[
 '1 รัน SELECT AVG(Salary) FROM Employees; แล้วบันทึกค่าที่ได้',
 '2 ใช้ Subquery ใน WHERE เพื่อหาพนักงานที่ Salary สูงกว่าค่าเฉลี่ยทั้งบริษัท ห้ามพิมพ์ค่าเฉลี่ยเป็นตัวเลขตายตัว',
 '3 ใช้ IN หรือ EXISTS เพื่อหาลูกค้าที่มีรายการสั่งซื้อ พร้อมอธิบายความหมายของคำสั่ง']),
 13:('สร้าง View และตรวจแผนการค้นหา','Case13_Performance_Script.sql',[
 '1 รันสคริปต์ตัวอย่าง แล้วสร้าง View สำหรับโครงงานของตนเองอย่างน้อย 1 ตัว ระบุว่าดึงข้อมูลจากตารางใดและใช้ทำอะไร',
 '2 เขียน SELECT อ่าน View พร้อมแนบผลลัพธ์และภาพ Navigator ที่เห็นชื่อ View',
 '3 เปรียบเทียบ EXPLAIN ก่อนและหลังมี INDEX อธิบายผลโดยไม่สรุปว่าข้อมูลเล็กต้องใช้ INDEX เสมอ']),
 14:('เขียน Stored Procedure รับพารามิเตอร์','Case14_Procedures_Triggers.sql',[
 '1 เขียน Procedure รับ ProductID 1 ค่าแล้วคืน ProductID, ProductName และ Price แนบ CREATE PROCEDURE และ CALL ที่ใช้ทดสอบ',
 '2 ทดลองทั้งรหัสที่มีและไม่มีใน Products แล้วอธิบายผลลัพธ์',
 '3 ตรวจ ProductDeleteLog หลังลบ ProductID 5 ในตัวอย่าง อธิบาย OLD และเวลาที่ Trigger ทำงาน']),
 15:('สร้างบัญชีอ่านยอดขายอย่างเดียว','Case15_Security_Guide.pdf',[
 "1 ใน MySQL ฝึกที่ครูจัดให้ เขียน CREATE USER 'report_admin'@'localhost' ใช้รหัสผ่าน 1234 ตามโจทย์ และให้ SELECT เฉพาะ db_city_case10.Sales",
 '2 ตรวจ SHOW GRANTS แล้วเชื่อมต่อด้วยบัญชีนี้ แสดงว่า SELECT ทำได้ แต่ INSERT UPDATE DELETE ถูกปฏิเสธ แนบผลโดยไม่เปิดเผยรหัสผ่านจริง',
 '3 เขียน REVOKE SELECT และอธิบายวิธีทดสอบหลังเชื่อมต่อใหม่'])}
for case,(title,starter,tasks) in worksheets.items():
    d=doc_base(f'ใบงาน {case:02d} {title}',case)
    ptext(d,f'ไฟล์ประกอบ {starter} เก็บสำเนาต้นฉบับไว้ก่อนเริ่ม และทำงานในฐานข้อมูลฝึกเฉพาะชุดนี้')
    if case==15:
        ptext(d,'ใช้ Case10_Aggregate_Data.sql เตรียมตาราง Sales ก่อน รหัส 1234 เป็นตัวอย่างในเครื่องฝึกเท่านั้น หากนโยบายรหัสผ่านปฏิเสธ ให้ใช้รหัสผ่านที่ครูกำหนดและบันทึกเหตุผล')
    for n,task in enumerate(tasks):
        if n>0 and case in [9,10] and n==3:d.add_page_break()
        answer(d,task,3)
    image_box(d,'แนบผลลัพธ์หรือภาพหลักฐานที่จำเป็น',4)
    save_doc(case,f'Worksheet_{case:02d}.docx',d,f'ใบงานปฏิบัติ {title}')
d=doc_base('แฟ้มสะสมผลงานฐานข้อมูลฉบับสมบูรณ์',16)
ptext(d,'ชื่อโครงงาน __________________________________________________________',True)
ptext(d,'ผู้จัดทำ _______________________ ชั้น __________ ครูผู้สอน __________________')
ptext(d,'รายวิชา 20901-2001 การจัดการฐานข้อมูลเบื้องต้น')
ptext(d,'ใช้โครงร่างนี้รวบรวมผลงานที่ทดสอบได้จริง แทนข้อความในวงเล็บด้วยข้อมูลโครงงานของตนเอง ไม่แนบรหัสผ่านจริง')
d.add_page_break();d.add_heading('1 ชื่อและวัตถุประสงค์ระบบงาน',1)
answer(d,'ชื่อระบบ ขอบเขต ผู้ใช้ และวัตถุประสงค์',5)
d.add_heading('2 พจนานุกรมข้อมูล',1)
ptext(d,'ทำซ้ำตารางต่อไปนี้สำหรับทุกตารางในระบบ ระบุ FK และตารางที่อ้างอิงในช่องข้อจำกัด')
t=d.add_table(rows=1,cols=5);t.style='Light Shading Accent 1'
for c,v in zip(t.rows[0].cells,['ตารางและฟิลด์','ชนิดข้อมูล','PK หรือ FK','NULL และข้อจำกัด','ความหมาย']):c.text=v
for _ in range(4):t.add_row()
d.add_page_break();d.add_heading('3 แผนภาพ EER และกฎธุรกิจ',1)
image_box(d,'วางภาพ EER Diagram ที่อ่านชื่อฟิลด์และความสัมพันธ์ได้',10)
answer(d,'อธิบายความสัมพันธ์และกฎธุรกิจ',3)
d.add_page_break();d.add_heading('4 สคริปต์ SQL ฉบับสมบูรณ์',1)
ptext(d,'วาง DDL และ DML ตามลำดับการรัน รวม SELECT, JOIN, Subquery, View, Index, Procedure และ Trigger ที่ใช้จริง ส่งไฟล์ .sql แยกควบคู่กัน')
answer(d,'ชื่อไฟล์ SQL และคำแนะนำการนำเข้า',3)
ptext(d,'[วางสคริปต์ SQL ที่นี่ เพิ่มหน้าตามความจำเป็น]')
d.add_heading('5 ผลทดสอบและคู่มือใช้งาน',1)
answer(d,'ระบุขั้นตอนทดสอบ ผลที่คาดหวัง ผลจริง และวิธีสำรองหรือกู้คืน',5)
save_doc(16,'Portfolio_Template.docx',d,'โครงร่าง Portfolio พร้อมหน้าปก Data Dictionary EER และ SQL')

# PDF pages are authored as HTML so Thai shaping and diagrams are preserved.
style='''@page{size:A4;margin:17mm}*{box-sizing:border-box}body{font-family:Tahoma,sans-serif;color:#172333;font-size:12px;line-height:1.65;margin:0}h1{font-size:23px;color:#000;line-height:1.4;margin:0 0 10px}h2{font-size:16px;color:#000;margin:15px 0 5px}p{margin:7px 0}header{font-size:10px;color:#516176;border-bottom:1px solid #cbd5e1;padding-bottom:8px;margin-bottom:14px}table{border-collapse:collapse;width:100%;margin:12px 0;table-layout:fixed}th,td{border:1px solid #d9d9d9;padding:9px;overflow-wrap:anywhere;vertical-align:middle}th{background:#163b56;color:white}tr:nth-child(even) td{background:#f3f7fb}pre{font-family:Consolas,monospace;white-space:pre-wrap;overflow-wrap:anywhere;background:#f1f5f9;border:1px solid #d9d9d9;padding:12px;font-size:11px;line-height:1.6}.note{background:#eef6fc;border-left:4px solid #2391bb;padding:10px}.shot{max-width:100%;max-height:470px;object-fit:contain;border:1px solid #ccc;display:block;margin:10px auto}.source{font-size:9px;color:#516176;overflow-wrap:anywhere}.page{break-after:page}.page:last-child{break-after:auto}svg{max-width:100%}li{margin-bottom:7px}'''
def page(title,body,case):
    week=case if case<=8 else case+1
    return f'<section class="page"><header>DATABASE CITY · รายวิชา 20901-2001 · หน่วย {case:02d} · สัปดาห์ {week}</header><h1>{title}</h1>{body}</section>'
pdf_jobs=[]
def pdf(case,filename,label,body):
    file=QA/(filename+'.html');file.write_text('<!doctype html><html lang="th"><meta charset="utf-8"><style>'+style+'</style><body>'+body+'</body></html>',encoding='utf-8')
    pdf_jobs.append({'html':str(file),'output':str(OUT/filename)});register(case,filename,label)

# Obtain genuine UI images from official MySQL documentation; never fabricate screenshots.
sources={}
def official_image(key,page_url,pattern):
    try:
        text=urllib.request.urlopen(page_url,timeout=20).read().decode('utf-8')
        imgs=re.findall(r'<img[^>]+src="([^"]+)"',text)
        candidates=[u for u in imgs if re.search(pattern,u,re.I)]
        if not candidates:return None
        url=urllib.parse.urljoin(page_url,candidates[0]);data=urllib.request.urlopen(url,timeout=20).read()
        target=QA/(key+'.png');target.write_bytes(data);sources[key]=(page_url,url);return target
    except Exception as e:
        print('Official image unavailable:',key,type(e).__name__);return None
root_image=official_image('root-password','https://dev.mysql.com/doc/mysql-installer/en/server-accounts-roles.html',r'accounts|roles')
connect_image=official_image('local-connection','https://dev.mysql.com/doc/workbench/en/wb-mysql-connections-new.html',r'new|connection')
reverse_image=official_image('reverse-engineer','https://dev.mysql.com/doc/workbench/en/wb-reverse-engineer-live.html',r'reverse-engineer-database-wizard')
setup_image=official_image('setup-type','https://dev.mysql.com/doc/refman/8.0/en/mysql-installer-setup.html',r'setup.type|choosing')
def fallback_image(key,url,source):
    target=QA/(key+'.png')
    try:
        data=urllib.request.urlopen(url,timeout=20).read();target.write_bytes(data)
        sources[key]=(source,url);return target
    except Exception as e:
        print('Screenshot unavailable:',key,type(e).__name__);return None
if root_image is None:
    root_image=fallback_image('root-password','https://www.tutorialsbook.com/layout/images/installmysql10.png','https://www.tutorialsbook.com/mysql/install-mysql')
if setup_image is None:
    setup_image=fallback_image('setup-type','https://kinsta.com/wp-content/uploads/2021/06/mysql-installer-option.jpg','https://kinsta.com/blog/mysql-community-server/')
if setup_image is None:
    setup_image=fallback_image('setup-type','https://terminalroot.com.br/assets/img/windows/php/chossing-a-setup-dev-default.jpg','https://terminalroot.com.br/2023/05/como-instalar-o-apache-php-8-e-mysql-no-windows.html')
def screenshot(path,caption,url):
    if path:
        source=sources.get(path.stem,(url,url))[0]
        return f'<img class="shot" src="{path.as_uri()}"><p class="source">ภาพหน้าจออ้างอิงเพื่อการเรียน ไม่ใช่ภาพการติดตั้งในเครื่องผู้เรียน · {html.escape(source)}</p><p>{caption}</p>'
    return f'<p class="note">{caption}</p><p class="source">ดูภาพหน้าจอจริงในคู่มือทางการ {html.escape(url)}</p>'

employee_table='''<table><tr><th>EmployeeID</th><th>FullName</th><th>Department</th><th>Salary</th></tr><tr><td>101</td><td>มาลี ใจดี</td><td>IT</td><td>28000</td></tr><tr><td>102</td><td>สมชาย สุขใจ</td><td>Sales</td><td>25000</td></tr><tr><td>103</td><td>มาลี ใจดี</td><td>HR</td><td>24000</td></tr></table>'''
infographic='''<svg viewBox="0 0 680 190" xmlns="http://www.w3.org/2000/svg"><defs><marker id="a" markerWidth="7" markerHeight="7" refX="6" refY="3" orient="auto"><path d="M0 0L6 3L0 6" fill="#15789b"/></marker></defs><rect x="7" y="8" width="664" height="171" rx="7" fill="#f5fbff" stroke="#15789b" stroke-width="3"/><text x="25" y="36" font-family="Tahoma" font-size="16">Employees</text><rect x="23" y="50" width="634" height="34" fill="#dff3dd" stroke="#318445" stroke-width="2"/><ellipse cx="100" cy="67" rx="70" ry="16" fill="none" stroke="#c67417" stroke-width="3"/><text x="45" y="73" font-family="Tahoma" font-size="14">EmployeeID</text><text x="215" y="73" font-family="Tahoma" font-size="14">FullName</text><text x="390" y="73" font-family="Tahoma" font-size="14">Department</text><text x="565" y="73" font-family="Tahoma" font-size="14">Salary</text><rect x="23" y="90" width="634" height="35" fill="#fff0e2" stroke="#b74c35" stroke-width="2"/><text x="60" y="113" font-size="14">101</text><text x="215" y="113" font-family="Tahoma" font-size="14">Mali Jaidee</text><text x="410" y="113" font-size="14">IT</text><text x="565" y="113" font-size="14">28000</text><path d="M55 154L15 154L15 120" fill="none" stroke="#15789b" stroke-width="2" marker-end="url(#a)"/><text x="65" y="160" font-family="Tahoma" font-size="13">Entity / Relation</text><text x="260" y="160" font-family="Tahoma" font-size="13" fill="#318445">Attribute / Field</text><text x="460" y="160" font-family="Tahoma" font-size="13" fill="#b74c35">Tuple / Record</text></svg>'''
pdf(2,'Case02_Table_Structure.pdf','ภาพสรุป Entity Attribute Tuple และ Primary Key',page('โครงสร้างตารางและคำศัพท์เชิงสัมพันธ์',f'<p>ตัวอย่าง Employees ใช้อธิบายเอนทิตีพนักงาน โดยแถวหนึ่งแทนพนักงานหนึ่งคน และคอลัมน์หนึ่งแทนคุณลักษณะหนึ่งชนิด</p>{infographic}{employee_table}<h2>อ่านกรอบสีและสัญลักษณ์</h2><ul><li>กรอบสีน้ำเงินรอบตาราง แทน Entity พนักงานในตัวอย่าง ตารางเชิงสัมพันธ์เรียก Relation</li><li>กรอบสีเขียวรอบหัวคอลัมน์ ชี้ชุด Attribute หรือ Field แต่ละชื่อคือหนึ่งคุณลักษณะ</li><li>กรอบสีส้มรอบแถว 101 คือ Tuple หรือ Record หนึ่งรายการ</li><li>วงรีสีทองที่ EmployeeID คือ Primary Key ใช้ระบุแถวโดยไม่ซ้ำและไม่เป็น NULL</li></ul><p class="note">FullName อาจซ้ำกันได้ จึงไม่เหมาะเป็นคีย์หลัก นักเรียนใช้แนวคิดนี้ตอบ Worksheet_02.docx</p>',2))
pdf(3,'Case03_Requirements.pdf','โจทย์กรณีศึกษาร้านกาแฟ Database City Coffee',page('โจทย์ออกแบบระบบร้านกาแฟ', '''<p>ร้าน Database City Coffee ต้องการบันทึกว่าใครขายเครื่องดื่มอะไร เมื่อใด และจำนวนเท่าใด เพื่อค้นหาประวัติการขายและสรุปยอดตามพนักงานได้</p><h2>ข้อมูลที่ต้องจัดเก็บ</h2><ul><li>Employees รหัสพนักงานและชื่อพนักงาน พนักงานหนึ่งคนทำรายการขายได้หลายรายการ</li><li>Drinks รหัสเครื่องดื่ม ชื่อเมนู และราคาปัจจุบัน เครื่องดื่มหนึ่งชนิดอยู่ได้ในหลายรายการขาย</li><li>Sales รหัสรายการขาย พนักงานผู้ขาย วันที่และเวลาขาย</li><li>SaleItems รหัสรายการขาย รหัสเครื่องดื่ม จำนวน และราคาต่อหน่วยขณะขาย</li></ul><h2>กฎธุรกิจ</h2><p>หนึ่งรายการขายอาจมีเครื่องดื่มหลายชนิด แต่ละชนิดอยู่หนึ่งบรรทัดต่อรายการขาย เก็บราคาขณะขายแยกจากราคาปัจจุบันเพื่อให้ประวัติไม่เปลี่ยนเมื่อร้านปรับราคา</p><p>ใช้คีย์ผสม SaleID และ DrinkID ใน SaleItems พนักงานและเมนูที่ถูกอ้างอิงต้องมีอยู่จริง กำหนดจำนวนมากกว่า 0 และราคาไม่ติดลบ</p><h2>ภารกิจ</h2><ol><li>หา Entity และ Attribute พร้อมกำหนด PK และ FK</li><li>วาด Chen ERD บน ERD_Template.pdf ระบุ 1 และ N ที่เส้นสัมพันธ์</li><li>อธิบายว่า Sales กับ Drinks เชื่อมกันผ่าน SaleItems อย่างไร</li><li>ส่งภาพ ERD ที่อ่านได้ พร้อมกฎธุรกิจอย่างน้อย 3 ข้อ</li></ol><p class="note">กรณีศึกษานี้ใช้ฝึกออกแบบ โครงงานต่อเนื่องของวิชายังคงเลือกระบบห้องสมุดหรือระบบที่ครูอนุมัติได้</p>''',3))
grid='<svg viewBox="0 0 680 780" xmlns="http://www.w3.org/2000/svg"><defs><pattern id="g" width="17" height="17" patternUnits="userSpaceOnUse"><path d="M17 0H0V17" fill="none" stroke="#d9e2e9" stroke-width=".6"/></pattern></defs><rect width="505" height="770" fill="url(#g)" stroke="#8093a1"/><rect x="525" y="10" width="145" height="755" fill="#f3f7fb" stroke="#cbd5e1"/><rect x="550" y="55" width="95" height="45" fill="white" stroke="#273c52"/><text x="570" y="84" font-size="14">Entity</text><ellipse cx="598" cy="175" rx="60" ry="30" fill="white" stroke="#273c52"/><text x="568" y="181" font-size="14">Attribute</text><path d="M598 245L660 290L598 335L535 290Z" fill="white" stroke="#273c52"/><text x="547" y="295" font-size="12">Relationship</text><text x="542" y="380" font-family="Tahoma" font-size="12">ขีดเส้นใต้ชื่อ PK</text><text x="542" y="420" font-family="Tahoma" font-size="12">ระบุ 1 หรือ N ที่เส้น</text><text x="542" y="460" font-family="Tahoma" font-size="12">ใช้ Chen notation</text><text x="542" y="500" font-family="Tahoma" font-size="12">เขียนชื่อความสัมพันธ์</text></svg>'
pdf(3,'ERD_Template.pdf','กระดาษกราฟวาด ERD พร้อมคำอธิบายสัญลักษณ์',page('กระดาษกราฟสำหรับออกแบบ ERD','<p>ชื่อ __________________________ รหัส _____________ ระบบ __________________</p>'+grid,3))
install=page('ติดตั้ง MySQL และเลือกส่วนประกอบ','''<p>คู่มือนี้ใช้หน้าจอ MySQL Installer สาย 8.0 บน Windows ชื่อเมนูอาจต่างตามรุ่น หากใช้ Server 8.4 หรือใหม่กว่าให้ติดตั้ง Server และ Workbench แยกตามชุดที่ครูกำหนด</p><h2>1 เลือก Setup Type</h2><p>ดาวน์โหลดจาก dev.mysql.com/downloads/installer/ เลือก Custom แล้วเลือก MySQL Server และ MySQL Workbench หากรุ่นที่ใช้มี Developer Default สามารถเลือกได้ ตรวจว่ามี Server อยู่ในรายการก่อน Execute</p>'''+screenshot(setup_image,'ตรวจส่วนประกอบให้มีทั้ง Server และ Workbench ก่อนเริ่มติดตั้ง','https://dev.mysql.com/doc/mysql-installer/en/mysql-installer-setup.html')+'<h2>2 ตั้งค่า Server</h2><p>เลือก Development Computer ใช้พอร์ตที่ครูกำหนด โดยค่าทั่วไปคือ 3306 ไม่เปิดพอร์ตสู่เครือข่ายสาธารณะสำหรับงานฝึกในเครื่อง ตรวจข้อกำหนดแล้วกด Execute</p>',4)
install+=page('ตั้งรหัสผ่านผู้ดูแล MySQL','<h2>3 Accounts and Roles</h2><p><strong>ห้ามลืม Root Password เด็ดขาด</strong> เก็บรหัสในที่ปลอดภัยและห้ามแนบในใบงาน ใช้รหัสของตนเอง ไม่ใช้รหัสผ่านจริงจากบริการอื่น</p>'+screenshot(root_image,'กรอก Root Password และ Repeat Password ให้ตรงกัน จากนั้นตั้ง Windows Service และ Apply Configuration จนรายการสำเร็จ','https://dev.mysql.com/doc/mysql-installer/en/server-accounts-roles.html'),4)
install+=page('สร้าง Local Connection ใน Workbench','<h2>4 เปิด MySQL Workbench</h2><p>กดเครื่องหมาย + ใกล้ MySQL Connections ตั้งชื่อ Local Classroom เลือก Standard TCP/IP กรอก Hostname 127.0.0.1, Port 3306 หรือตามที่ตั้งไว้ และ Username ตามบัญชีที่ครูจัดให้</p>'+screenshot(connect_image,'กด Test Connection ใส่รหัสผ่านเมื่อถาม แล้วเปิด Connection ที่สร้าง','https://dev.mysql.com/doc/workbench/en/wb-mysql-connections-new.html')+'<h2>5 ตรวจว่าเชื่อมต่อได้จริง</h2><pre>SELECT VERSION(), CURRENT_USER();</pre><p>ต้องมีแถวผลลัพธ์และไม่เกิดข้อความ error จากนั้นบันทึกภาพลง Worksheet_04.docx ถ้าเชื่อมต่อไม่ได้ตรวจบริการ MySQL พอร์ต และบัญชีอีกครั้ง</p>',4)
pdf(4,'Installation_Guide_MySQL.pdf','คู่มือติดตั้งและเชื่อมต่อ MySQL พร้อมภาพหน้าจออ้างอิง',install)
pdf(5,'Worksheet_05.pdf','สรุปไวยากรณ์ DDL สำหรับใช้คู่กับแม่แบบ SQL',page('สรุปคำสั่งสร้างฐานข้อมูลและตาราง','''<p>ใช้ทบทวนระหว่างทำ Case05_DDL_Template.sql เปลี่ยนชื่อฐานข้อมูลหากมีงานเก่าอยู่แล้ว</p><h2>สร้างและเลือกฐานข้อมูล</h2><pre>CREATE DATABASE db_city CHARACTER SET utf8mb4;
USE db_city;</pre><h2>สร้างตาราง</h2><pre>CREATE TABLE tb_users (
  id INT PRIMARY KEY,
  username VARCHAR(50) NOT NULL
);</pre><h2>ตรวจโครงสร้าง</h2><pre>SHOW TABLES;
DESCRIBE tb_users;
SHOW CREATE TABLE tb_users;</pre><h2>ลบตารางฝึกเมื่อจำเป็น</h2><pre>-- ตรวจชื่อ schema และสำรองก่อนลบ
-- DROP TABLE tb_users;</pre><p class="note">DROP TABLE ลบทั้งโครงสร้างและข้อมูล จึงแสดงเป็นคอมเมนต์ ไม่ต้องรันระหว่างตรวจคำตอบตามปกติ</p><h2>งานที่ส่ง</h2><p>ไฟล์ .sql ที่เติมครบ พร้อมภาพ SHOW TABLES และ DESCRIBE ที่สำเร็จ</p>''',5))
pdf(6,'Case06_DataTypes_Reference.pdf','สรุปชนิดข้อมูล MySQL พร้อม PK FK',page('ชนิดข้อมูลและข้อจำกัดใน MySQL','''<table><tr><th>หมวด</th><th>ชนิดข้อมูล</th><th>การใช้งาน</th></tr><tr><td>ตัวเลข</td><td>INT</td><td>รหัสและจำนวนเต็ม</td></tr><tr><td>ตัวเลข</td><td>DECIMAL(M,D)</td><td>ราคาที่ต้องการทศนิยมแน่นอน M คือจำนวนหลักทั้งหมด D คือจำนวนหลักทศนิยม เช่น DECIMAL(10,2)</td></tr><tr><td>ตัวอักษร</td><td>VARCHAR(N)</td><td>ข้อความความยาวแปรผันสูงสุด N ตัวอักษร เหมาะกับชื่อและเบอร์โทร</td></tr><tr><td>ตัวอักษร</td><td>TEXT</td><td>ข้อความยาว เช่น รายละเอียดสินค้า</td></tr><tr><td>วันเวลา</td><td>DATE</td><td>วันที่ เช่น 2026-11-24</td></tr><tr><td>วันเวลา</td><td>DATETIME</td><td>วันและเวลา เช่น 2026-11-24 09:00:00</td></tr></table><h2>Primary Key และ Foreign Key</h2><pre>PRIMARY KEY (CustomerID)
FOREIGN KEY (CustomerID)
  REFERENCES Customers(CustomerID)</pre><p>PK ต้องไม่ซ้ำและไม่เป็น NULL ส่วน FK ป้องกันการอ้างอิงข้อมูลที่ไม่มีจริง ใช้ InnoDB และชนิดคอลัมน์ที่เข้ากัน</p><p class="note">Case06_Script_Template.sql เป็นแม่แบบที่มีช่อง ___ ต้องเติมก่อนรัน ทดลองเพิ่ม Order ที่อ้าง CustomerID ที่ไม่มีอยู่ แล้วบันทึกข้อผิดพลาดที่ได้</p>''',6))
eer=page('สร้าง EER จากฐานข้อมูลฝึก','<p>รัน Case08_Master_Script.sql ก่อน ชุดนี้สร้าง 4 ตารางและข้อมูลครบ แต่ยังไม่ประกาศ FK เพื่อให้ฝึกเชื่อมเส้นเอง</p><h2>1 Reverse Engineer</h2><p>เปิดเมนู Database แล้วเลือก Reverse Engineer เลือก Local Connection กด Next และเลือก schema db_city_case08 เลือกทั้ง 4 ตารางและ Place imported objects on a diagram แล้ว Finish</p>'+screenshot(reverse_image,'ภาพ wizard จากคู่มือทางการ เมนูอาจต่างเล็กน้อยตามรุ่น','https://dev.mysql.com/doc/workbench/en/wb-reverse-engineer-live.html'),8)
eer+=page('เชื่อมความสัมพันธ์และตรวจ Foreign Key','''<h2>2 เพิ่มเส้นในแบบจำลอง</h2><p>เลือกเครื่องมือ 1:N non-identifying relationship กดตารางลูกก่อนแล้วกดตารางแม่ ตรวจหน้า Foreign Keys ว่าใช้คอลัมน์เดิม ไม่สร้างคอลัมน์ซ้ำโดยไม่ตั้งใจ</p><ul><li>Orders.CustomerID อ้าง Customers.CustomerID</li><li>OrderItems.OrderID อ้าง Orders.OrderID</li><li>OrderItems.ProductID อ้าง Products.ProductID</li></ul><h2>อ่านชนิดความสัมพันธ์</h2><table><tr><th>ชนิด</th><th>ความหมาย</th></tr><tr><td>1 ต่อ 1</td><td>แถวหนึ่งเชื่อมได้ไม่เกินหนึ่งแถวอีกฝั่ง มักต้องมี UNIQUE ที่ FK</td></tr><tr><td>1 ต่อ N</td><td>ลูกค้าหนึ่งคนมีหลายใบสั่งซื้อ FK อยู่ฝั่งหลาย</td></tr><tr><td>M ต่อ N</td><td>Orders และ Products เชื่อมผ่าน OrderItems แยกเป็น 1:N สองเส้น</td></tr></table><p>เส้นทึบหรือเส้นประใน Workbench แสดง identifying หรือ non-identifying relationship ไม่ได้ใช้แทนจำนวน 1 หรือ N โดยตรง</p><h2>3 ส่งมอบและตรวจผล</h2><p>บันทึกไฟล์ .mwb และภาพ EER การแก้แบบจำลองยังไม่เปลี่ยน Server จนกว่าจะ Forward Engineer หรือรัน ALTER TABLE ที่ตรวจแล้วใน schema ฝึก จากนั้นตรวจ SHOW CREATE TABLE และทดสอบ FK</p>''',8)
pdf(8,'Case08_EER_Guide.pdf','คู่มือ Reverse Engineer และสร้างความสัมพันธ์พร้อมภาพอ้างอิง',eer)
security=page('จัดการบัญชีและสิทธิ์อ่านยอดขาย','''<p>ใช้บัญชีผู้ดูแลที่ครูจัดให้เฉพาะ MySQL ฝึก รัน Case10_Aggregate_Data.sql เพื่อเตรียม db_city_case10.Sales ก่อนเริ่ม</p><h2>1 สร้างบัญชีตามโจทย์</h2><pre>CREATE USER 'report_admin'@'localhost'
  IDENTIFIED BY '1234';</pre><p>1234 เป็นรหัสตัวอย่างในห้องทดลองเท่านั้น หาก password policy ปฏิเสธ ให้ใช้รหัสผ่านที่ครูกำหนดแทนโดยไม่ปิดนโยบายระบบ</p><h2>2 ให้เฉพาะสิทธิ์ SELECT ตาราง Sales</h2><pre>GRANT SELECT ON db_city_case10.Sales
  TO 'report_admin'@'localhost';
SHOW GRANTS FOR 'report_admin'@'localhost';</pre><h2>3 เข้าใจ FLUSH PRIVILEGES</h2><pre>-- FLUSH PRIVILEGES;</pre><p>CREATE USER, GRANT และ REVOKE มีผลโดยไม่ต้อง FLUSH PRIVILEGES คำสั่งนี้เกี่ยวกับการโหลด grant tables ใหม่ เช่น หลังแก้ grant tables โดยตรง ซึ่งไม่ใช่วิธีที่ใช้ในใบงานนี้</p><h2>4 ทดสอบและถอนสิทธิ์</h2><pre>-- เปิด Connection ใหม่เป็น report_admin แล้วลอง
SELECT * FROM db_city_case10.Sales;
-- INSERT UPDATE DELETE ต้องถูกปฏิเสธ

-- กลับบัญชีผู้ดูแล
REVOKE SELECT ON db_city_case10.Sales
  FROM 'report_admin'@'localhost';</pre><p>ตรวจ SHOW GRANTS อีกครั้ง และเชื่อมต่อใหม่ก่อนทดสอบสิทธิ์หลังถอน ห้ามรันการแก้สิทธิ์กับฐานข้อมูลจริง</p><p class="source">อ้างอิง MySQL Reference Manual https://dev.mysql.com/doc/refman/8.0/en/privilege-changes.html</p>''',15)
pdf(15,'Case15_Security_Guide.pdf','คู่มือ CREATE USER GRANT FLUSH PRIVILEGES และ REVOKE',security)
rubric='''<p>คะแนนรวม 100 คะแนน ใช้กับโครงงานฐานข้อมูลที่ครูอนุมัติ ผู้เรียนต้องส่งไฟล์ SQL, EER, Portfolio และแสดงผลทดสอบได้</p><table><tr><th style="width:21%">ด้านประเมิน</th><th style="width:10%">เต็ม</th><th>รายละเอียดคะแนนย่อยและหลักฐาน</th></tr><tr><td>โครงสร้างตาราง</td><td>20</td><td>Entity และขอบเขตเหมาะสม 5 คะแนน ชนิดข้อมูลเหมาะสม 5 คะแนน PK และข้อจำกัด 5 คะแนน นำเข้า DDL ได้จริง 5 คะแนน</td></tr><tr><td>ความสัมพันธ์และ FK</td><td>20</td><td>EER ตรงกับตาราง 5 คะแนน cardinality ถูกต้อง 5 คะแนน FK ทำงานจริง 5 คะแนน ทดสอบรหัสอ้างอิงผิดและอธิบายผล 5 คะแนน</td></tr><tr><td>Query View และ Trigger</td><td>40</td><td>SELECT และเงื่อนไข 8 คะแนน JOIN 8 คะแนน Subquery 6 คะแนน View 6 คะแนน Procedure และ Trigger 8 คะแนน INDEX และ EXPLAIN 4 คะแนน</td></tr><tr><td>เอกสารและการส่งมอบ</td><td>20</td><td>วัตถุประสงค์และ Data Dictionary 5 คะแนน EER อ่านได้ 5 คะแนน SQL และคู่มือนำเข้า 5 คะแนน ผลทดสอบและสำรองกู้คืน 5 คะแนน</td></tr></table><h2>ระดับคุณภาพของแต่ละรายการย่อย</h2><ul><li>ครบคะแนน ทำงานได้ถูกต้อง หลักฐานครบ และอธิบายเหตุผลได้ด้วยตนเอง</li><li>ประมาณ 75 เปอร์เซ็นต์ ทำงานหลักได้ มีข้อผิดพลาดเล็กน้อยและแก้ไขได้</li><li>ประมาณ 50 เปอร์เซ็นต์ ทำงานได้บางส่วน หลักฐานหรือคำอธิบายยังไม่ครบ</li><li>ประมาณ 25 เปอร์เซ็นต์ มีเพียงโครงร่าง ยังแสดงการทำงานจริงไม่ได้</li><li>0 คะแนน ไม่ส่งหรือไม่มีหลักฐานในรายการนั้น</li></ul><p>ครูให้คะแนนย่อยตามหลักฐานจริงในช่วง 0 ถึงคะแนนเต็มของแต่ละรายการ และรวมได้ไม่เกิน 100 คะแนน</p>'''
pdf(16,'Final_Project_Rubric.pdf','รูบริกประเมินโครงงานฉบับสมบูรณ์ 100 คะแนน',page('เกณฑ์ประเมินโครงงานปลายภาค',rubric,16))

from structured_worksheets import build_worksheets
from workbench_resources import build_datasets
build_datasets(sql, setup, OUT)
build_worksheets(doc_base,ptext,set_font,save_doc,manifest)
(QA/'pdf-jobs.json').write_text(json.dumps(pdf_jobs,ensure_ascii=False,indent=2),encoding='utf-8')
(QA/'resource-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
(QA/'image-sources.json').write_text(json.dumps(sources,ensure_ascii=False,indent=2),encoding='utf-8')
print(f'Created {len(list(OUT.glob("*.sql")))} SQL files and {len(list(OUT.glob("*.docx")))} Word documents; queued {len(pdf_jobs)} PDFs')
